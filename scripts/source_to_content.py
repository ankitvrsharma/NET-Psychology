#!/usr/bin/env python3
"""Repository source extraction/cache and deterministic application of validated manual ChatGPT content."""
from pathlib import Path
import argparse, hashlib, json, os, re, subprocess, sys
from copy import deepcopy
from datetime import datetime, timezone
import tempfile

ROOT=Path(__file__).resolve().parents[1]
INBOX=ROOT/"sources/inbox"
EXTRACTION_CACHE=ROOT/".cache/source-extraction"
EXTRACTOR_VERSION="3"
SUPPORTED_SOURCE_SUFFIXES={".pdf",".docx",".pptx",".txt",".md"}
SYLLABUS=ROOT/"data/syllabus-index.json"
QUICK=ROOT/"content/quick-learn/quick_cards.json"
QUESTIONS=ROOT/"content/questions/questions.json"
REVISION=ROOT/"content/revision/revision_guidance.json"
MICRO=ROOT/"content/microtopics/micro_topics.json"
STAGING=ROOT/"content-staging"
VERIFICATION=ROOT/"data/verification-state.json"
VERIFICATION_KEYS={"microtopic":"microtopics","deep_dive":"deepDive","active_recall":"activeRecall","revision":"revision","practice":"practice"}

def load(p): return json.loads(Path(p).read_text(encoding="utf-8"))
def save(p,o):
    Path(p).parent.mkdir(parents=True,exist_ok=True)
    Path(p).write_text(json.dumps(o,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
def norm(s): return re.sub(r"[^a-z0-9]+"," ",str(s).lower()).strip()
def sha(p):
    h=hashlib.sha256()
    with open(p,"rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""): h.update(b)
    return h.hexdigest()

def _ocr_pdf_page(path,page_number):
    """OCR one low-text PDF page when Poppler and Tesseract are available."""
    try:
        with tempfile.TemporaryDirectory(prefix="net-psychology-ocr-") as tmp:
            prefix=Path(tmp)/"page"
            subprocess.run(
                ["pdftoppm","-f",str(page_number),"-l",str(page_number),"-r","180","-png","-singlefile",str(path),str(prefix)],
                capture_output=True,text=True,check=True,timeout=120
            )
            image=prefix.with_suffix(".png")
            if not image.exists():
                return ""
            result=subprocess.run(
                ["tesseract",str(image),"stdout","-l",os.getenv("NET_OCR_LANG","eng"),"--psm","3"],
                capture_output=True,text=True,check=True,timeout=120
            )
            return result.stdout.strip()
    except (OSError,subprocess.SubprocessError):
        return ""

def _useful_text_length(text):
    return len(re.sub(r"\s+","",str(text or "")))

def extract(p):
    """Extract text once, retaining PDF page boundaries and OCRing low-text pages."""
    if p.suffix.lower()==".pdf":
        pages=[]
        try:
            import pypdf
            reader=pypdf.PdfReader(str(p))
            pages=[(page.extract_text() or "").strip() for page in reader.pages]
        except Exception:
            pages=[]
        # Poppler sometimes extracts more usable text than the PDF text layer.
        if pages and sum(_useful_text_length(x) for x in pages)<200:
            try:
                fallback=subprocess.run(
                    ["pdftotext","-layout",str(p),"-"],
                    capture_output=True,text=True,check=True,timeout=180
                ).stdout
                fallback_pages=fallback.split("\f")
                if fallback_pages and not fallback_pages[-1].strip():
                    fallback_pages.pop()
                if sum(_useful_text_length(x) for x in fallback_pages)>sum(_useful_text_length(x) for x in pages):
                    pages=[x.strip() for x in fallback_pages]
            except (OSError,subprocess.SubprocessError):
                pass
        if not pages:
            try:
                page_count=int(subprocess.run(
                    ["pdfinfo",str(p)],capture_output=True,text=True,check=True,timeout=30
                ).stdout.split("Pages:",1)[1].splitlines()[0].strip())
                pages=[""]*page_count
            except (OSError,subprocess.SubprocessError,IndexError,ValueError):
                return ""
        output=[]
        for index,page_text in enumerate(pages,1):
            # OCR only pages whose text layer is missing or nearly empty.
            if _useful_text_length(page_text)<60:
                ocr_text=_ocr_pdf_page(p,index)
                if _useful_text_length(ocr_text)>_useful_text_length(page_text):
                    page_text=ocr_text
            output.append(f"[PAGE {index}]\n{page_text.strip()}")
        return "\n\n".join(output)
    if p.suffix.lower() in {".txt",".md"}:
        return p.read_text(encoding="utf-8",errors="ignore")
    if p.suffix.lower()==".docx":
        from docx import Document
        return "\n".join(x.text for x in Document(str(p)).paragraphs)
    if p.suffix.lower()==".pptx":
        from pptx import Presentation
        return "\n".join(sh.text for slide in Presentation(str(p)).slides for sh in slide.shapes if hasattr(sh,"text"))
    return ""

def _cache_file(cache_root,relative_path,source_hash):
    key=hashlib.sha256(f"{EXTRACTOR_VERSION}\0{relative_path}\0{source_hash}".encode("utf-8")).hexdigest()
    return Path(cache_root)/(key+".json")

def _read_cached_extraction(cache_file,relative_path,source_hash):
    try:
        entry=json.loads(Path(cache_file).read_text(encoding="utf-8"))
        if (entry.get("extractor_version")==EXTRACTOR_VERSION
            and entry.get("source_path")==relative_path
            and entry.get("source_sha256")==source_hash
            and isinstance(entry.get("text"),str)):
            return entry["text"]
    except (OSError,ValueError,TypeError):
        pass
    return None

def cached_extract(path,relative_path,cache_root=EXTRACTION_CACHE):
    """Return (text, cache_hit); cache key binds source path, bytes and extractor version."""
    source_hash=sha(path)
    cache_file=_cache_file(cache_root,relative_path,source_hash)
    cached=_read_cached_extraction(cache_file,relative_path,source_hash)
    if cached is not None:
        return cached,True
    text=extract(path)
    entry={
        "extractor_version":EXTRACTOR_VERSION,
        "source_path":relative_path,
        "source_sha256":source_hash,
        "characters":len(text),
        "extracted_at":datetime.now(timezone.utc).isoformat(),
        "text":text
    }
    cache_file.parent.mkdir(parents=True,exist_ok=True)
    temp_path=cache_file.with_suffix(".tmp")
    temp_path.write_text(json.dumps(entry,ensure_ascii=False)+"\n",encoding="utf-8")
    os.replace(temp_path,cache_file)
    return text,False

def load_source_library(inbox=INBOX,cache_root=EXTRACTION_CACHE):
    """Load cached source extractions, re-extracting only changed or missing entries."""
    inbox=Path(inbox)
    files=sorted(p for p in inbox.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SOURCE_SUFFIXES)
    manifest={str(p.relative_to(ROOT)).replace("\\","/"):sha(p) for p in files}
    cache_root=Path(cache_root)
    cache_root.mkdir(parents=True,exist_ok=True)
    # Drop stale cache records for changed/deleted sources; never keep orphaned text.
    for cache_file in cache_root.glob("*.json"):
        try:
            entry=json.loads(cache_file.read_text(encoding="utf-8"))
            rel=entry.get("source_path")
            if rel not in manifest or entry.get("source_sha256")!=manifest.get(rel) or entry.get("extractor_version")!=EXTRACTOR_VERSION:
                cache_file.unlink(missing_ok=True)
        except (OSError,ValueError,TypeError):
            cache_file.unlink(missing_ok=True)
    source_chunks=[]; source_meta=[]; cache_hits=0; cache_misses=0
    for path in files:
        rel=str(path.relative_to(ROOT)).replace("\\","/")
        text,hit=cached_extract(path,rel,cache_root)
        cache_hits+=int(hit); cache_misses+=int(not hit)
        pieces=chunks(text)
        source_meta.append({
            "path":rel,"sha256":manifest[rel],"chunks":len(pieces),
            "characters":len(text),"extraction_cache":"hit" if hit else "miss"
        })
        source_chunks.extend((rel,index,chunk) for index,chunk in enumerate(pieces))
    stats={"files":len(source_meta),"chunks":len(source_chunks),"characters":sum(x["characters"] for x in source_meta),
           "cache_hits":cache_hits,"cache_misses":cache_misses}
    return source_chunks,source_meta,stats

def chunks(text,size=12000):
    text=re.sub(r"\n{3,}","\n\n",text).strip()
    out=[]; cur=""
    for part in text.split("\n\n"):
        if len(cur)+len(part)+2<=size: cur += ("\n\n" if cur else "")+part
        else:
            if cur: out.append(cur)
            cur=part[:size]
    if cur: out.append(cur)
    return out

def canonical(s):
    out={}
    for u in s["units"]:
        for t in u.get("topics",[]):
            for m in t.get("microtopics",[]):
                k=f"{u['id']}-{t['id']}-{m['id']}"
                out[k]={"id":k,"unit":u["id"],"topic":t["id"],"micro":m["id"],
                        "title":m.get("title",""),"topic_title":t.get("title",""),"unit_title":u.get("title","")}
    return out

def candidate_refs(text,refs,n=8):
    words=set(norm(text).split()); ranked=[]
    for k,r in refs.items():
        ts=set(norm(r["title"]+" "+r["topic_title"]+" "+r["unit_title"]).split())
        ranked.append((len(words&ts),r))
    ranked.sort(key=lambda x:x[0],reverse=True)
    return [r for _,r in ranked[:n]]

def main():
    parser=argparse.ArgumentParser(description="Extract/cache approved repository sources or apply validated ChatGPT content.")
    parser.add_argument("--prepare-source-cache",action="store_true",help="Extract/cache repository sources; no content generation.")
    parser.add_argument("--apply-staged",action="store_true",help="Apply the validated manual ChatGPT response.")
    parser.add_argument("--apply",action="store_true",help=argparse.SUPPRESS)
    parser.add_argument("--repair-existing",action="store_true",help=argparse.SUPPRESS)
    parser.add_argument("--enrich-existing",action="store_true",help=argparse.SUPPRESS)
    parser.add_argument("--rewrite-existing",action="store_true",help=argparse.SUPPRESS)
    parser.add_argument("--source-triggered",action="store_true",help=argparse.SUPPRESS)
    args=parser.parse_args()
    if args.apply or args.repair_existing or args.enrich_existing or args.rewrite_existing or args.source_triggered:
        raise SystemExit("The legacy model-generation path has been retired. Use scripts/prepare_chatgpt_packet.py and the manual ChatGPT response workflow.")
    if args.prepare_source_cache:
        _,_,stats=load_source_library()
        print("Source extraction cache prepared: "+json.dumps(stats,sort_keys=True))
        return 0
    if not args.apply_staged:
        raise SystemExit("Choose --prepare-source-cache or --apply-staged. Automated model generation is no longer supported.")
    _,source_meta,stats=load_source_library()
    approved_path=STAGING/"canonical-content-approved.json"
    if not approved_path.exists(): raise SystemExit("Missing validated content-staging/canonical-content-approved.json.")
    staged=load(approved_path)
    if isinstance(staged,list):
        packages=staged; quick_updates=[]; question_updates=[]; operation="package_rewrite"
    elif isinstance(staged,dict):
        packages=staged.get("packages") or []; quick_updates=staged.get("quick_cards") or []
        question_updates=staged.get("question_updates") or []; operation=staged.get("operation","package_rewrite")
    else: raise SystemExit("Validated staged content has an invalid root shape.")
    micro_pool=load(MICRO); deep_pool=load(ROOT/"content/deep-dive/deep_dive.json")
    recall_pool=load(ROOT/"content/active-recall/active_recall.json"); revision_pool=load(REVISION)
    practice_path=ROOT/"content/practice/practice_mcqs.json"
    practice_pool=load(practice_path) if practice_path.exists() else {}
    verification=load(VERIFICATION) if VERIFICATION.exists() else {"schema_version":1,"items":{},"updated_at":""}
    verification.setdefault("items",{})
    applied={"packages":0,"quick_cards":0,"question_updates":0}
    if quick_updates:
        quick_pool=load(QUICK)
        if not isinstance(quick_pool,dict): raise SystemExit("Quick Learn pool must be an object keyed by card ID.")
        for row in quick_updates:
            rid=str(row["id"]); current=quick_pool.get(rid)
            if not isinstance(current,dict): raise SystemExit("Quick Learn card disappeared after validation: "+rid)
            current.update(row["updates"]); applied["quick_cards"]+=1
        save(QUICK,quick_pool)
    if question_updates:
        qstore=load(QUESTIONS)
        groups=list(qstore.values()) if isinstance(qstore,dict) else [qstore]
        for row in question_updates:
            rid=str(row["id"]); current=None
            for group in groups:
                if isinstance(group,list):
                    current=next((q for q in group if isinstance(q,dict) and str(q.get("id",""))==rid),None)
                    if current: break
            if current is None: raise SystemExit("Question disappeared after validation: "+rid)
            for field,value in row["updates"].items():
                actual=field
                if field=="question" and field not in current and "q" in current: actual="q"
                if field=="options" and field not in current and "o" in current: actual="o"
                current[actual]=value
            applied["question_updates"]+=1
        save(QUESTIONS,qstore)
    if packages:
        def locked(component,mid):
            return str(verification.setdefault("items",{}).setdefault(VERIFICATION_KEYS[component],{}).get(mid,""))=="EXPERT VERIFIED"
        for item in packages:
            mid=str(item.get("microtopic_id",""))
            if mid not in micro_pool: raise SystemExit("Canonical micro-topic disappeared after validation: "+mid)
            micro=micro_pool[mid]
            if not locked("microtopic",mid):
                micro["expert_explanation"]=item["core_explanation"]
                micro["cross_references"]=item.get("cross_references") or []
            if not locked("deep_dive",mid):
                deep=deep_pool.get(mid) or {"id":mid,"title":micro.get("title","")}
                deep.update(id=mid+"D",microtopic_id=mid,title=micro.get("title",deep.get("title","")),detailed_explanation=item["detailed_explanation"])
                deep_pool[mid]=deep
            if not locked("active_recall",mid):
                recall=recall_pool.get(mid) or {"id":mid,"title":micro.get("title",""),"prompts":[]}
                recall.update(id=mid+"A",microtopic_id=mid,title=micro.get("title",recall.get("title","")),prompts=item.get("recall_prompts") or [])
                recall_pool[mid]=recall
            if not locked("revision",mid):
                guidance=item.get("revision_guidance") or {}; revision=revision_pool.get(mid) or {"id":mid,"title":micro.get("title","")}
                revision.update(id=mid+"R",microtopic_id=mid,title=micro.get("title",revision.get("title","")))
                for field in ("recall_before_review","self_check","weak_point_prompt","rating_instruction"): revision[field]=guidance.get(field,"")
                revision_pool[mid]=revision
            if not locked("practice",mid) and item.get("practice_mcqs"):
                practice_pool[mid]={"id":mid+"P","microtopic_id":mid,"title":micro.get("title",""),"questions":item.get("practice_mcqs") or []}
            applied["packages"]+=1
        save(MICRO,micro_pool); save(ROOT/"content/deep-dive/deep_dive.json",deep_pool)
        save(ROOT/"content/active-recall/active_recall.json",recall_pool); save(REVISION,revision_pool)
        save(practice_path,practice_pool); save(VERIFICATION,verification)
    now=datetime.now(timezone.utc).isoformat()
    provenance_path=ROOT/"content-provenance.json"
    provenance=load(provenance_path) if provenance_path.exists() else {"schema_version":1,"sources":[],"runs":[]}
    provenance["generated_by"]="manual ChatGPT packet workflow"
    provenance["generated_at"]=now
    provenance["sources"]=source_meta
    provenance.setdefault("runs",[]).append({"operation":operation,"provider":"ChatGPT conversation (manual; no API)","generated_at":now,"sources":source_meta,"source_cache_stats":stats,"applied":applied})
    save(provenance_path,provenance)
    print("Applied validated ChatGPT content: "+json.dumps(applied,sort_keys=True))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
